"""tables.py - gom kết quả các lần chạy thành bảng và ghi results.xlsx (GUIDE.md mục 6.1)."""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd


def collect_summaries(out_dir: str | Path, prefix: str | None = None) -> pd.DataFrame:
    """Đọc mọi <out_dir>/<exp_id>/seed<k>/summary.json thành một DataFrame (một dòng mỗi lần chạy)."""
    rows = []
    for f in sorted(Path(out_dir).glob("*/seed*/summary.json")):
        r = json.loads(f.read_text())
        cfg = json.loads((f.parent / "config.json").read_text())
        r.update({k: cfg[k] for k in ("img_size", "aug", "loss", "mix", "sampler", "ema_decay", "lr_backbone",
                                      "lr_head", "label_smoothing", "class_weight_beta", "batch_size")})
        rows.append(r)
    df = pd.DataFrame(rows)
    if prefix and len(df):
        df = df[df["exp_id"].str.startswith(prefix)]
    return df.reset_index(drop=True)


def mean_std_table(df: pd.DataFrame, by: str = "exp_id", cols=("val_macro_f1", "val_top1")) -> pd.DataFrame:
    """mean ± std (ddof=1) theo seed cho mỗi exp_id; std là NaN nếu chỉ có 1 seed."""
    g = df.groupby(by)
    out = g[list(cols)].agg(["mean", lambda s: s.std(ddof=1)])
    out.columns = [f"{c}_{'mean' if m == 'mean' else 'std'}" for c, m in out.columns]
    out["n_seeds"] = g.size()
    return out.reset_index()


def write_results_xlsx(path: str | Path, sheets: dict[str, pd.DataFrame], highlight: dict[str, str] | None = None,
                       digits: int = 4) -> Path:
    """Ghi các sheet bằng openpyxl: freeze hàng tiêu đề, số thập phân thống nhất, tô dòng tốt nhất.

    highlight: {tên_sheet: tên_cột} -> tô dòng có giá trị lớn nhất của cột đó.
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for j, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(f"{v:.{digits}f}" if isinstance(v, float) else str(v))
                                              for v in df[col].head(200)])
                ws.column_dimensions[get_column_letter(j)].width = min(45, width + 2)
                if pd.api.types.is_float_dtype(df[col]):
                    for cell in ws[get_column_letter(j)][1:]:
                        cell.number_format = "0." + "0" * digits
            col = (highlight or {}).get(name)
            if col in df.columns and df[col].notna().any():
                r = int(np.nanargmax(df[col].to_numpy(dtype=float))) + 2
                for cell in ws[r]:
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
    return path
